from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SHOT_DIR = DOCS / "codecanyon-screenshots"
OUT = DOCS / "Vertue_CRM_CodeCanyon_Feature_Documentation.docx"


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded(draw, box, fill, outline=None, width=1, radius=16):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw, xy, value, size=24, fill="#0f172a", bold=False, max_width=None, line_gap=6):
    f = font(size, bold)
    if max_width:
        avg = max(size * 0.52, 7)
        lines = []
        for paragraph in str(value).split("\n"):
            lines.extend(wrap(paragraph, max(8, int(max_width / avg))) or [""])
        x, y = xy
        for line in lines:
            draw.text((x, y), line, font=f, fill=fill)
            y += size + line_gap
        return y
    draw.text(xy, str(value), font=f, fill=fill)
    return xy[1] + size


def draw_sidebar(draw, active):
    rounded(draw, (22, 22, 238, 878), "#082f49", radius=22)
    text(draw, (42, 45), "Vertue CRM", 26, "#e0f2fe", True)
    items = ["Dashboard", "Students", "Universities", "Applications", "Tasks", "Finance", "Settings"]
    y = 100
    for item in items:
        fill = "#0ea5e9" if item == active else "#0b3d5c"
        rounded(draw, (42, y, 218, y + 44), fill, radius=12)
        text(draw, (58, y + 12), item, 17, "#ffffff" if item == active else "#cfe8f8", True)
        y += 56


def screenshot_dashboard(path):
    img = Image.new("RGB", (1400, 900), "#eef6fb")
    d = ImageDraw.Draw(img)
    draw_sidebar(d, "Dashboard")
    text(d, (280, 45), "CRM Workspace Dashboard", 34, "#0f172a", True)
    text(d, (280, 85), "Admissions KPIs, pipeline, revenue, automation and notifications", 18, "#64748b")
    cards = [
        ("Students", "1,248", "#dbeafe"),
        ("Active Applications", "386", "#dcfce7"),
        ("Pending Tasks", "42", "#fef3c7"),
        ("New Requests", "17", "#fae8ff"),
    ]
    x = 280
    for title, value, color in cards:
        rounded(d, (x, 130, x + 285, 245), "#ffffff", "#dbe4ee", radius=16)
        rounded(d, (x + 18, 150, x + 66, 198), color, radius=12)
        text(d, (x + 84, 150), title, 18, "#64748b", True)
        text(d, (x + 84, 184), value, 34, "#0f172a", True)
        x += 310
    rounded(d, (280, 285, 820, 600), "#ffffff", "#dbe4ee", radius=18)
    text(d, (310, 310), "Pipeline Funnel", 24, "#0f172a", True)
    for i, (label, pct, col) in enumerate([("Lead", 64, "#38bdf8"), ("Applied", 42, "#22c55e"), ("Enrolled", 18, "#f59e0b")]):
        y = 370 + i * 65
        text(d, (315, y), label, 18, "#334155", True)
        rounded(d, (430, y, 765, y + 24), "#e2e8f0", radius=12)
        rounded(d, (430, y, 430 + int(335 * pct / 100), y + 24), col, radius=12)
        text(d, (780, y - 2), f"{pct}%", 17, "#0f172a", True)
    rounded(d, (850, 285, 1310, 600), "#ffffff", "#dbe4ee", radius=18)
    text(d, (880, 310), "Upcoming Tasks", 24, "#0f172a", True)
    for i, task in enumerate(["Follow up missing documents", "Visa call with student", "University offer review", "Payment confirmation"]):
        y = 365 + i * 52
        rounded(d, (880, y, 1280, y + 38), "#f8fafc", "#e2e8f0", radius=10)
        text(d, (900, y + 10), task, 16, "#334155")
    rounded(d, (280, 640, 1310, 850), "#ffffff", "#dbe4ee", radius=18)
    text(d, (310, 670), "Revenue Forecast + Lead Source ROI", 24, "#0f172a", True)
    for i, h in enumerate([70, 130, 95, 160, 125, 185]):
        x = 340 + i * 105
        rounded(d, (x, 820 - h, x + 52, 820), "#0ea5e9", radius=10)
    text(d, (980, 710), "Conversion Rate: 21.4%\nVisa Success: 76.0%\nCAC: $94.50", 22, "#0f172a", True)
    img.save(path)


def screenshot_students(path):
    img = Image.new("RGB", (1400, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    draw_sidebar(d, "Students")
    text(d, (280, 45), "Student Management", 34, "#0f172a", True)
    rounded(d, (280, 105, 1310, 170), "#ffffff", "#dbe4ee", radius=16)
    filters = ["Search name/email/phone", "Stage", "Country", "GPA min", "GPA max", "50 per page"]
    x = 305
    for f in filters:
        rounded(d, (x, 122, x + 150, 154), "#f1f5f9", "#cbd5e1", radius=9)
        text(d, (x + 10, 131), f, 13, "#475569")
        x += 160
    headers = ["Student", "Nationality", "GPA", "Field", "Agent", "Sub-Agent", "Stage", "Actions"]
    y = 220
    rounded(d, (280, y, 1310, 820), "#ffffff", "#dbe4ee", radius=18)
    x_positions = [305, 475, 610, 700, 860, 990, 1130, 1240]
    for x, h in zip(x_positions, headers):
        text(d, (x, y + 25), h, 15, "#64748b", True)
    rows = [
        ("Aylin Demir", "Turkey", "3.4", "Computer Science", "Agent A", "Sub 1", "Applicant"),
        ("Omar Khalid", "Jordan", "3.1", "Business", "Agent B", "Sub 2", "Visa"),
        ("Sara Ahmadi", "Iran", "3.7", "Medicine", "Agent A", "-", "Enrolled"),
        ("Mina Kaya", "Turkey", "2.9", "Architecture", "Agent C", "Sub 3", "Lead"),
        ("John Smith", "USA", "3.2", "Law", "Agent B", "-", "Docs Pending"),
    ]
    y += 70
    for row in rows:
        d.line((300, y - 12, 1290, y - 12), fill="#e2e8f0", width=1)
        for x, value in zip(x_positions, row + ("View Edit",)):
            color = "#0ea5e9" if value in ["Applicant", "Visa", "Enrolled", "Lead", "Docs Pending"] else "#334155"
            text(d, (x, y), value, 15, color, value in ["Applicant", "Visa", "Enrolled", "Lead", "Docs Pending"])
        y += 78
    text(d, (300, 790), "Includes documents, portal accounts, password reset, ownership scope and student activity timeline.", 17, "#64748b")
    img.save(path)


def screenshot_universities(path):
    img = Image.new("RGB", (1400, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    draw_sidebar(d, "Universities")
    text(d, (280, 45), "University Catalog + Programs", 34, "#0f172a", True)
    for i, (label, val) in enumerate([("Programs Total", "3,420"), ("Missing Fee / Language", "42 / 18"), ("Missing Thesis Type", "12")]):
        x = 280 + i * 345
        rounded(d, (x, 105, x + 320, 190), "#ffffff", "#dbe4ee", radius=16)
        text(d, (x + 20, 125), label, 16, "#64748b", True)
        text(d, (x + 20, 152), val, 26, "#0f172a", True)
    rounded(d, (280, 220, 1310, 290), "#ffffff", "#dbe4ee", radius=16)
    text(d, (305, 242), "Search | Country | Date From | Date To | Sort | 50/100 per page | Grid/List | Import | Dedupe | Delete All", 17, "#334155")
    cards = [
        ("Istanbul Technical University", "Turkey / Istanbul", "Bachelor - Computer Engineering", "$3,500 yearly"),
        ("Near East University", "Northern Cyprus / Nicosia", "Medicine, Dentistry, Business", "$6,000 per semester"),
        ("Ankara Science University", "Turkey / Ankara", "Associate, Bachelor, Master", "$2,800 yearly"),
    ]
    x = 280
    for title, meta, programs, fee in cards:
        rounded(d, (x, 330, x + 320, 690), "#ffffff", "#dbe4ee", radius=18)
        rounded(d, (x + 20, 350, x + 300, 470), "#dbeafe", radius=14)
        text(d, (x + 20, 495), title, 20, "#0f172a", True, 270)
        text(d, (x + 20, 548), meta, 15, "#64748b")
        text(d, (x + 20, 585), programs, 15, "#334155", False, 270)
        text(d, (x + 20, 645), fee, 16, "#0ea5e9", True)
        x += 345
    rounded(d, (280, 725, 1310, 835), "#ecfeff", "#bae6fd", radius=18)
    text(d, (310, 755), "Built-in duplicate cleanup removes repeated universities, repeated programs and duplicate applications while preserving the most complete row.", 22, "#0f172a", True, 980)
    img.save(path)


def screenshot_applications(path):
    img = Image.new("RGB", (1400, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    draw_sidebar(d, "Applications")
    text(d, (280, 45), "Applications Desk", 34, "#0f172a", True)
    rounded(d, (280, 110, 1310, 185), "#ffffff", "#dbe4ee", radius=16)
    text(d, (305, 135), "Dynamic university selector by country, program selector, intake, status, deadline and follow-up date", 18, "#334155")
    rounded(d, (280, 225, 1310, 805), "#ffffff", "#dbe4ee", radius=18)
    headers = ["ID", "Student", "University", "Program", "Status", "Enroll %", "Next Action"]
    xs = [305, 400, 570, 790, 980, 1110, 1210]
    for x, h in zip(xs, headers):
        text(d, (x, 255), h, 15, "#64748b", True)
    rows = [
        ("#1052", "Aylin Demir", "Istanbul Tech", "Computer Engineering", "Documents Pending", "72%", "Collect transcript"),
        ("#1053", "Omar Khalid", "Near East", "Medicine", "Visa Process", "84%", "Visa fee follow-up"),
        ("#1054", "Sara Ahmadi", "Ankara Science", "Business", "Offer Sent", "65%", "Confirm offer"),
    ]
    y = 315
    for row in rows:
        d.line((300, y - 14, 1290, y - 14), fill="#e2e8f0", width=1)
        for x, value in zip(xs, row):
            text(d, (x, y), value, 15, "#334155", value.endswith("%") or value.startswith("#"))
        y += 88
    rounded(d, (320, 620, 1265, 755), "#f0fdf4", "#bbf7d0", radius=18)
    text(d, (350, 650), "Duplicate protection", 24, "#166534", True)
    text(d, (350, 690), "The system prevents duplicated applications for the same student, university, program and intake.", 18, "#166534", False, 860)
    img.save(path)


def screenshot_portal(path):
    img = Image.new("RGB", (1200, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    rounded(d, (35, 25, 1165, 875), "#ffffff", "#dbe4ee", radius=26)
    text(d, (75, 65), "Student Portal", 34, "#0f172a", True)
    text(d, (75, 105), "Dashboard, applications, document progress, university matching and messaging", 18, "#64748b")
    rounded(d, (75, 165, 505, 315), "#e0f2fe", radius=18)
    text(d, (105, 195), "Document Progress", 22, "#075985", True)
    rounded(d, (105, 255, 455, 278), "#bae6fd", radius=12)
    rounded(d, (105, 255, 360, 278), "#0284c7", radius=12)
    text(d, (470, 248), "73%", 28, "#075985", True)
    rounded(d, (545, 165, 1125, 315), "#f0fdf4", radius=18)
    text(d, (575, 195), "My Applications", 22, "#166534", True)
    text(d, (575, 240), "Offer Sent\nVisa Process\nDocuments Pending", 18, "#166534")
    rounded(d, (75, 355, 1125, 695), "#ffffff", "#e2e8f0", radius=18)
    text(d, (105, 385), "Recommended Universities", 24, "#0f172a", True)
    for i, title in enumerate(["Istanbul Technical University", "Near East University", "Ankara Science University"]):
        y = 450 + i * 78
        rounded(d, (105, y, 1085, y + 55), "#f8fafc", "#e2e8f0", radius=12)
        text(d, (130, y + 15), title, 17, "#334155", True)
        text(d, (770, y + 15), "Apply / Send Request", 15, "#0ea5e9", True)
    rounded(d, (95, 755, 1105, 830), "#0f172a", radius=22)
    for i, item in enumerate(["Home", "Apps", "Docs", "Unis", "More"]):
        x = 145 + i * 195
        color = "#38bdf8" if i == 0 else "#e2e8f0"
        text(d, (x, 782), item, 18, color, True)
    img.save(path)


def screenshot_saas(path):
    img = Image.new("RGB", (1400, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    draw_sidebar(d, "Settings")
    text(d, (280, 45), "SaaS Super Admin", 34, "#0f172a", True)
    rounded(d, (280, 105, 1310, 255), "#ffffff", "#dbe4ee", radius=18)
    text(d, (310, 135), "Tenant Manager", 24, "#0f172a", True)
    text(d, (310, 175), "Create tenant, assign subscription, generate subdomain, approve account, enable/disable modules.", 18, "#334155")
    rounded(d, (280, 300, 790, 790), "#ffffff", "#dbe4ee", radius=18)
    text(d, (310, 330), "SaaS Packages", 24, "#0f172a", True)
    for i, p in enumerate(["Starter Monthly", "Growth Quarterly", "Scale Semiannual", "Enterprise Annual"]):
        y = 385 + i * 82
        rounded(d, (310, y, 750, y + 58), "#f8fafc", "#e2e8f0", radius=12)
        text(d, (330, y + 15), p, 18, "#334155", True)
        text(d, (620, y + 15), "$", 18, "#0ea5e9", True)
    rounded(d, (830, 300, 1310, 790), "#ffffff", "#dbe4ee", radius=18)
    text(d, (860, 330), "Tenant Access", 24, "#0f172a", True)
    text(d, (860, 385), "vista.virtuevisa.com\ncustom-domain support\nsubscription status\nfeature toggles\nuser activation", 21, "#334155", False, 420)
    img.save(path)


def screenshot_mobile(path):
    img = Image.new("RGB", (760, 1100), "#e2e8f0")
    d = ImageDraw.Draw(img)
    rounded(d, (170, 35, 590, 1065), "#0f172a", radius=48)
    rounded(d, (190, 75, 570, 1025), "#f8fafc", radius=34)
    text(d, (225, 115), "CRM Mobile", 30, "#0f172a", True)
    rounded(d, (225, 175, 535, 260), "#ffffff", "#e2e8f0", radius=18)
    text(d, (245, 200), "Students", 18, "#64748b", True)
    text(d, (245, 228), "1,248", 28, "#0f172a", True)
    for i, label in enumerate(["Pipeline", "Tasks", "Applications", "Universities", "Finance"]):
        y = 300 + i * 92
        rounded(d, (225, y, 535, y + 64), "#ffffff", "#e2e8f0", radius=16)
        text(d, (250, y + 20), label, 20, "#334155", True)
    rounded(d, (205, 930, 555, 1000), "#ffffff", "#cbd5e1", radius=24)
    for i, item in enumerate(["Home", "Apps", "Tasks", "More"]):
        x = 230 + i * 82
        text(d, (x, 955), item, 14, "#0ea5e9" if i == 0 else "#64748b", True)
    img.save(path)


def add_heading(doc, text_value, level=1):
    p = doc.add_heading(text_value, level=level)
    return p


def add_bullets(doc, items, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        doc.add_paragraph(item, style=style)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    return table


def set_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 24), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        styles[name].font.name = "Calibri"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(15, 23, 42)


def build_doc():
    DOCS.mkdir(exist_ok=True)
    SHOT_DIR.mkdir(exist_ok=True)
    screenshots = {
        "dashboard": SHOT_DIR / "01-dashboard.png",
        "students": SHOT_DIR / "02-students.png",
        "universities": SHOT_DIR / "03-universities.png",
        "applications": SHOT_DIR / "04-applications.png",
        "portal": SHOT_DIR / "05-student-portal.png",
        "saas": SHOT_DIR / "06-saas.png",
        "mobile": SHOT_DIR / "07-mobile.png",
    }
    screenshot_dashboard(screenshots["dashboard"])
    screenshot_students(screenshots["students"])
    screenshot_universities(screenshots["universities"])
    screenshot_applications(screenshots["applications"])
    screenshot_portal(screenshots["portal"])
    screenshot_saas(screenshots["saas"])
    screenshot_mobile(screenshots["mobile"])

    doc = Document()
    set_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Vertue CRM - Multi-Tenant Education Consultancy CRM")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(14, 116, 144)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Complete Codecanyon Feature Documentation and Product Overview").bold = True
    doc.add_paragraph(
        "This document describes every major feature included in the Vertue CRM source code. It is written for buyers, reviewers and marketplace listing preparation, with module-by-module capability details, SaaS functionality, portal features, security architecture and product preview screenshots."
    )
    doc.add_picture(str(screenshots["dashboard"]), width=Inches(7.2))

    add_heading(doc, "1. Product Summary", 1)
    doc.add_paragraph(
        "Vertue CRM is a Laravel-based, multi-tenant admissions CRM and student portal built for education agencies, visa consultants, university recruitment offices and international student admission teams. The system combines lead management, student records, university/program catalogs, application tracking, task management, finance, reports, messaging, automation, SaaS tenant management and a student self-service portal."
    )
    add_table(doc, ["Item", "Details"], [
        ("Technology Stack", "Laravel 10.x, PHP 8.1+, MySQL, Blade templates, vanilla JavaScript, responsive CSS"),
        ("Architecture", "Multi-tenant CRM with tenant isolation, SaaS subscriptions, tenant feature toggles and role/permission based access"),
        ("Primary Users", "Super Admin, Tenant Admin, Agent, Sub-Agent, Student Portal User"),
        ("Main Use Case", "Manage international student recruitment from lead intake to application, documents, finance, visa process and enrollment"),
        ("Commercial Fit", "Education consultancy CRM, SaaS CRM, admissions office system, student portal platform"),
    ])

    add_heading(doc, "2. Highlight Features for Marketplace Listing", 1)
    add_bullets(doc, [
        "Multi-tenant SaaS architecture with tenant subscription status, package plans, feature toggles, subdomain/custom-domain fields and tenant activation workflow.",
        "Complete student CRM with stages, agents, sub-agents, GPA, nationality, study field, target country, lead source, lifecycle stage and portal login accounts.",
        "University and program catalog with bulk CSV import, export template, country/date filters, grid/list view, program fee type, duplicate university/program cleanup and scholarship links.",
        "Application desk with dynamic university/program selectors, intake tracking, admission statuses, deadline, follow-up date, scoring, explainability and duplicate application protection.",
        "Student portal with dashboard, university matching, application requests, document progress, offer letters and secure messaging.",
        "Task management with priority, deadlines, completion, creator/assignee visibility and permission-aware task scoping.",
        "Advanced search across students, universities, programs and applications with degree, thesis, language, country, city and export options.",
        "Finance module with payments, commissions, status filtering, multi-currency settings and SaaS sales summaries for super admins.",
        "Reports, audit logs, automation rules, templates, API tokens, webhook endpoint, notifications and backup/restore tooling.",
        "Mobile-friendly CRM layout with compact pagination, responsive tables, slide-out sidebar and bottom navigation for CRM and student portal users.",
    ])

    add_heading(doc, "3. Screenshots and Product Preview", 1)
    for label, key in [
        ("Dashboard and KPI Overview", "dashboard"),
        ("Student List and Filters", "students"),
        ("University Catalog and Program Manager", "universities"),
        ("Applications Desk", "applications"),
        ("Student Portal", "portal"),
        ("SaaS Tenant Management", "saas"),
        ("Mobile Bottom Navigation", "mobile"),
    ]:
        add_heading(doc, label, 2)
        doc.add_picture(str(screenshots[key]), width=Inches(6.7 if key == "mobile" else 7.2))

    add_heading(doc, "4. User Roles and Access Control", 1)
    add_table(doc, ["Role", "Typical Access"], [
        ("Super Admin", "SaaS tenants, packages, tenant subscriptions, all CRM modules, settings and system-level management."),
        ("Admin", "Tenant-wide CRM management depending on permissions: students, universities, applications, finance, tasks, reports, templates and settings."),
        ("Agent", "Assigned students, assigned sub-agents, scoped applications/tasks/messages and student follow-up workflows."),
        ("Sub-Agent", "Scoped student/application/task access based on assignment and permissions."),
        ("Student", "Student portal dashboard, application status, document view, university matching, requests and messages."),
    ])
    doc.add_paragraph(
        "The CRM uses permission middleware for module access and additional ownership scoping for sensitive data. Admins can be configured with global access while agents and sub-agents can be restricted to their own students, tasks and related records."
    )
    add_bullets(doc, [
        "Role permission manager supports module-level view/create/update/delete access.",
        "Scope permissions support own-data vs global-data access, including students, universities, applications, finance, messages, scholarships, requests and tasks.",
        "Task visibility is secured so assignees see tasks sent to them by the correct sender scope; admins with view-all permission can see all tenant tasks.",
        "Tenant isolation middleware prevents cross-tenant data access.",
        "Subscription middleware blocks inactive, suspended or expired tenant accounts.",
    ])

    add_heading(doc, "5. SaaS Platform Module", 1)
    add_bullets(doc, [
        "Create, edit, approve and manage SaaS tenants.",
        "Tenant records include company name, owner name, email, phone, currency, plan type, subscription start/end dates, subscription status and login access.",
        "Subdomain, custom domain and access URL fields are included for tenant-specific access.",
        "SaaS packages include name, slug, price, currency, duration, feature list, active status and sort order.",
        "Default package examples include Starter Monthly, Growth Quarterly, Scale Semiannual and Enterprise Annual.",
        "Feature toggles can enable or disable modules per tenant, including applications, messaging, reports export, task management, multi-currency, file upload, sub-agent creation, WhatsApp notifications, API tokens, automation rules, backup, mobile navigation and duplicate cleanup tools.",
        "Super admin finance view includes SaaS sales, active subscriptions and expired subscriptions.",
    ])

    add_heading(doc, "6. Dashboard Module", 1)
    add_bullets(doc, [
        "Dashboard cards for total students, active applications, pending tasks and new student requests.",
        "Pipeline counts for lead, applied and enrolled stages.",
        "Recent student activity and unread notifications.",
        "Top programs report based on application volume.",
        "Upcoming and overdue tasks.",
        "Monthly revenue chart data and next-month revenue forecast.",
        "Conversion rate, visa success rate, lead-source ROI and customer acquisition cost calculations for users with reporting access.",
        "Data automatically respects tenant scope and user ownership rules.",
    ])

    add_heading(doc, "7. Student Management", 1)
    add_bullets(doc, [
        "Student list with search by name, email, phone and field of study.",
        "Filters for stage, target country, GPA minimum and GPA maximum.",
        "15/50/100 pagination menu with 50 records as the default list size.",
        "Student profile fields include full name, email, phone, nationality, GPA, field of study, preferred university language, English level, stage, lifecycle stage, target country, lead source, budget, passport number, agent and sub-agent.",
        "Create, update, delete and reset portal password.",
        "Student detail screen includes applications, documents, required-document checklist, tasks, messages, offer letters and AI insight fields when available.",
        "Document upload, view, verify and delete actions for staff users.",
        "Automatic portal user linkage through the student user account.",
        "WhatsApp notification hooks for new students and document updates when the tenant feature is enabled.",
    ])

    add_heading(doc, "8. Pipeline Board", 1)
    add_bullets(doc, [
        "Kanban-style student pipeline organized by student stages.",
        "Move students between stages through a pipeline move endpoint.",
        "User-specific pipeline preferences can be saved.",
        "Supports lead, inquiry, applicant, documents pending, interview scheduled, admitted, visa process, tuition paid, enrolled and alumni workflows.",
    ])

    add_heading(doc, "9. University and Program Catalog", 1)
    add_bullets(doc, [
        "University list with search, country filter, date range filter, sorting and 15/50/100 pagination.",
        "Grid/list display controls and compact icon toolbar.",
        "University fields include name, country, city, institution type, website, currency, tuition range, tuition fee type, language, deadline, visa notes, description, logo and image.",
        "Program manager supports multiple rows per university with program name, degree level, thesis type, language, duration, currency, fee, fee type and notes.",
        "Bulk CSV import and exportable CSV template.",
        "Duplicate cleanup tool removes duplicate universities, duplicate programs and duplicate applications.",
        "Program duplicate detection now ignores fee type so the same program does not appear repeatedly just because fee type differs.",
        "Study catalog synchronization can derive study fields from imported program names.",
        "Bulk delete, delete all, individual delete and program delete actions are included.",
        "Mojibake/encoding cleanup helpers improve imported catalog text readability.",
    ])

    add_heading(doc, "10. Applications Desk", 1)
    add_bullets(doc, [
        "Application list with search, status filter, university sorting and pagination.",
        "Create and edit applications for selected student, university, program, intake and status.",
        "Dynamic university options by country.",
        "Dynamic program options by selected university, degree, thesis type and language.",
        "Tracks deadline, next follow-up date and notes.",
        "Enrollment probability, explainability and best next action are calculated during create/update.",
        "Duplicate application guard prevents the same student from having repeated applications for the same university, program and intake.",
        "WhatsApp notifications can be triggered for application changes.",
    ])

    add_heading(doc, "11. Student Requests", 1)
    add_bullets(doc, [
        "Incoming portal application requests from students.",
        "Pending and processed tabs.",
        "Approve request and optionally assign an agent.",
        "Reject request flow.",
        "When approved, the request can create/convert student and application data and notify relevant staff.",
        "Ownership scope is applied for agents and sub-agents.",
    ])

    add_heading(doc, "12. Task Management", 1)
    add_bullets(doc, [
        "Task list filtered by status, priority and assignee for users with view-all access.",
        "Task fields include student, assigned user, creator, priority, status, deadline, description and escalation level.",
        "Create, update, delete and mark complete actions.",
        "Security-aware visibility: recipients only see tasks assigned to them or sent by the appropriate creator unless they have admin/view-all permission.",
        "Overdue tasks feed dashboard alerts and automation rules.",
    ])

    add_heading(doc, "13. Messaging and Notifications", 1)
    add_bullets(doc, [
        "CRM messaging module for staff-to-student communication.",
        "Student portal messaging with optional attachments.",
        "Notification center with unread counts, open links and mark-read/mark-all-read actions.",
        "Notification events are used for portal requests, student messages, automation follow-ups and SLA alerts.",
    ])

    add_heading(doc, "14. Scholarships Module", 1)
    add_bullets(doc, [
        "Scholarship list with search, country, university, date range and sort filters.",
        "Create, update and delete scholarships.",
        "Scholarship fields include university, title, discount percentage and description.",
        "University mapping is displayed in the scholarship table.",
        "Pagination supports 15/50/100 records.",
    ])

    add_heading(doc, "15. Advanced Search", 1)
    add_bullets(doc, [
        "Search across students, applications and universities from one advanced interface.",
        "Filters include keywords, university country, city, university type, university name, selected university, program name, degree, thesis type, program language, study field, student stage, student country, application status and preferred university language.",
        "Degree aliases normalize Diploma, Associate, Bachelor, Master and PhD labels.",
        "Program options are dynamically loaded by selected university and filters.",
        "Export options for matching universities and programs.",
    ])

    add_heading(doc, "16. Finance Module", 1)
    add_bullets(doc, [
        "Finance dashboard cards for total amount, paid amount, commission and outstanding value.",
        "Payment records include student, type, amount, currency, commission amount, status and paid date.",
        "Payment filters by status and currency.",
        "Create, edit and delete payments.",
        "Super admins see SaaS sales summaries and recent SaaS subscriptions.",
        "Multi-currency support through tenant currency settings.",
    ])

    add_heading(doc, "17. Reports and Analytics", 1)
    add_bullets(doc, [
        "Advanced student report filters by date range, agent and country.",
        "CSV export.",
        "Excel-compatible export.",
        "PDF export when Dompdf is installed, with HTML fallback when Dompdf is not available.",
        "Agent performance report with student, application, task, payment and sub-agent metrics.",
        "Dashboard analytics include conversion rate, visa success rate, revenue forecast and lead-source ROI.",
    ])

    add_heading(doc, "18. Automation Rules", 1)
    add_bullets(doc, [
        "Automation rule manager for tenant-level rules.",
        "Supported trigger keys include SLA overdue tasks, daily follow-up and documents pending follow-up configuration.",
        "SLA overdue task automation creates notifications and can trigger WhatsApp alerts.",
        "Documents pending automation can create follow-up tasks, notifications, WhatsApp messages, email and SMS depending on tenant settings.",
        "Rules can be created, activated/deactivated and manually run.",
    ])

    add_heading(doc, "19. Templates, API Tokens and Webhooks", 1)
    add_bullets(doc, [
        "Message template manager for reusable communication templates.",
        "API token management for tenant integrations.",
        "Webhook endpoint for student status updates.",
        "Integration settings include email, SMS and AI provider configuration fields.",
    ])

    add_heading(doc, "20. Settings Module", 1)
    add_bullets(doc, [
        "User profile settings: name, language, font scale, preferred currency and theme.",
        "Password change with current-password validation.",
        "Theme options include Figma light, Figma dark, classic light and classic dark.",
        "Currency manager with default currency support.",
        "WhatsApp notification provider settings including custom and Twilio-style provider fields.",
        "Documents-pending automation settings: number of days, task, email, WhatsApp and SMS toggles.",
        "Email, SMS and AI integration settings.",
        "Lead source cost settings for ROI and CAC reporting.",
    ])

    add_heading(doc, "21. Study Catalogs", 1)
    add_bullets(doc, [
        "Study field catalog management.",
        "Intake term catalog management.",
        "Sync study fields from university programs.",
        "Bulk delete and delete-all actions for catalog cleanup.",
        "Invalid/mojibake imported field cleanup support.",
        "Used by student records, search filters and portal university matching.",
    ])

    add_heading(doc, "22. Backup, Restore and Health Checks", 1)
    add_bullets(doc, [
        "Health screen checks database availability, storage availability, app environment, debug state and app URL.",
        "Tenant backup export generates JSON with metadata, counts and tenant-owned CRM data.",
        "Backup includes students, applications, documents, payments, tasks, messages, requests, universities, programs, scholarships, automation rules, templates, notifications, API tokens, study fields and intake terms when tables exist.",
        "Restore supports uploaded JSON file or pasted JSON payload and restores tenant-scoped rows safely.",
    ])

    add_heading(doc, "23. Student Portal", 1)
    add_bullets(doc, [
        "Separate student login and authentication guard.",
        "Student dashboard showing applications, required document progress and recent messages.",
        "University browser with filters for keywords, country, city, type, university name, degree and study field.",
        "University matching service ranks universities based on student profile, field, country, budget, language and GPA.",
        "Students can submit university application requests for admin review.",
        "Students can view applications and documents, including offer/acceptance letters.",
        "Students can send messages with attachments to staff.",
        "Mobile portal bottom navigation provides quick access to Home, Apps, Docs, Universities and More.",
    ])

    add_heading(doc, "24. Security and Tenant Isolation", 1)
    add_bullets(doc, [
        "Separate CRM and student authentication guards.",
        "JWT API authentication middleware for API endpoints.",
        "Tenant isolation via tenant ID on all main business records.",
        "Subscription active middleware blocks suspended or expired tenants.",
        "Role and permission middleware protects routes.",
        "Agent and sub-agent ownership scoping protects student/application/task visibility.",
        "Subdomain/custom-domain tenant access support with fallback to base domain during DNS setup.",
        "Audit logging for sensitive operations.",
    ])

    add_heading(doc, "25. Mobile and Responsive UI", 1)
    add_bullets(doc, [
        "Responsive CRM layout with slide-out sidebar on smaller screens.",
        "Bottom navigation for admin/staff CRM users.",
        "Bottom navigation for student portal users.",
        "Compact pagination with 15/50/100 record selection.",
        "Mobile-friendly tables with horizontal scroll instead of broken layout.",
        "Responsive forms, dialogs and toolbar controls.",
    ])

    add_heading(doc, "26. Included Pages and Modules", 1)
    add_table(doc, ["Area", "Pages / Functions"], [
        ("Public Website", "Landing page, register, privacy policy, contact, SaaS plans API"),
        ("Authentication", "CRM login/logout, student portal login/logout"),
        ("CRM", "Dashboard, students, pipeline, universities, applications, tasks, messages, scholarships, finance, search"),
        ("Administration", "Agents & roles, permissions, user-specific permissions, templates, API tokens, automation rules, audit logs"),
        ("SaaS", "Tenants, tenant profile, package plans, feature toggles, subscriptions, subdomain/access URL"),
        ("Student Portal", "Dashboard, universities, applications, documents, messages, apply-to-university request"),
        ("System", "Health checks, tenant backup, restore, notifications, webhooks"),
    ])

    add_heading(doc, "27. Buyer-Facing Summary", 1)
    doc.add_paragraph(
        "Vertue CRM is suitable for a buyer who needs a production-ready starting point for an education consultancy SaaS platform. It already includes the CRM core, student portal, multi-tenant SaaS management, subscription controls, permissions, reports, automation, backups, university/program catalog management and mobile-friendly navigation. The source is structured around Laravel controllers, Blade views and tenant-scoped database tables, making it extendable for custom branding, payment gateway integration, external university APIs, email/SMS providers and advanced AI features."
    )

    add_heading(doc, "28. Technical Requirements", 1)
    add_bullets(doc, [
        "PHP 8.1 or newer.",
        "Laravel Framework 10.x.",
        "MySQL/MariaDB database.",
        "Composer for dependency installation.",
        "Writable storage and public storage symlink for uploaded documents and university media.",
        "Optional Dompdf package for native PDF export; otherwise HTML fallback is available.",
        "DNS wildcard and SSL wildcard certificate are recommended for tenant subdomains.",
    ])

    add_heading(doc, "29. Important Notes for Marketplace Buyers", 1)
    add_bullets(doc, [
        "The application is multi-tenant; each tenant must have an active subscription record to access CRM pages.",
        "For SaaS subdomains, configure DNS such as *.yourdomain.com pointing to the server IP and set SAAS_BASE_DOMAIN in the environment file.",
        "Run all migrations or the latest SQL patch before using registration, SaaS tenants, feature toggles and duplicate cleanup modules.",
        "Configure APP_URL, database credentials, mail/SMS/WhatsApp providers and storage permissions before production deployment.",
        "Change demo credentials after installation and never use default credentials in production.",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
